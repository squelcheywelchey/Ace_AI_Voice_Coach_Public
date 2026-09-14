"""Seed calibration.db from the round-1 exports.

Reads:
  - config.LABELS_XLSX       coach_human_labels export (sheet "Human labels")
  - config.TRANSCRIPTS_XLSX  grading_platform_input_final.xlsx (row_idx -> excerpt)

Loads every transcript in the grading file (except the Tutorial row), freezes
all DONE round-1 labels into prior_labels (tagged G1-G4 via groups.py), creates
the two round-2 labelers, and assigns every transcript to both. Transcripts
whose round-1 grading was never finished simply have no prior labels.
Idempotent: wipes and reloads (including any round-2 scores!).

Run:  /usr/bin/python3 seed.py     (needs openpyxl)
"""

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook

import config
import db
from groups import group_of
# The ROUND-1 dimension ids — the columns of the labels xlsx. Fixed here rather
# than derived from rubric.py, which evolves (e.g. feedback_is_correct was later
# split into four bracket dims that don't exist in the round-1 data).
DIM_IDS = [
    "scaffolds_then_fades",
    "adapts_when_stuck",
    "kind_delivery",
    "reentry_appropriate_framing",
    "limits_the_load",
    "feedback_is_correct",
    "makes_it_a_dialogue",
    "drives_practice_uptake",
    "quality_conversational_flow",
    "pii",
]


def _assignments_table_exists() -> bool:
    try:
        db._query_one("SELECT 1 FROM assignments LIMIT 1")
        return True
    except Exception:
        return False


def norm_row(v) -> str:
    """source_row / row_idx values arrive as int or str — normalize."""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def read_transcripts():
    """{source_row: {ceo_id, transcript, graders}} for every non-Tutorial row."""
    ws = load_workbook(config.TRANSCRIPTS_XLSX, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    hdr = {h: i for i, h in enumerate(next(rows))}
    out = {}
    for r in rows:
        if r[hdr["row_idx"]] is None:
            continue
        src = norm_row(r[hdr["row_idx"]])
        entry = out.setdefault(src, {
            "ceo_id": norm_row(r[hdr["ceo_id"]]) if r[hdr["ceo_id"]] is not None else "",
            "transcript": str(r[hdr["excerpt"]] or ""),
            "graders": set(),
        })
        entry["graders"].add(str(r[hdr["grader_name"]] or ""))
    # drop rows that exist only as the tutorial example
    return {src: e for src, e in out.items() if e["graders"] != {"Tutorial"}}


def read_labels():
    ws = load_workbook(config.LABELS_XLSX, read_only=True, data_only=True)["Human labels"]
    rows = ws.iter_rows(values_only=True)
    hdr = {h: i for i, h in enumerate(next(rows))}
    labels = []
    for r in rows:
        if not r[hdr["labeler"]] or (r[hdr["status"]] or "") != "done":
            continue
        get = lambda col: (str(r[hdr[col]]).strip() if hdr.get(col) is not None and r[hdr[col]] is not None else "")
        scores = {d: get(d) for d in DIM_IDS if get(d)}
        notes = {d: get(f"note_{d}") for d in DIM_IDS if get(f"note_{d}")}
        evidence = {}
        for d in DIM_IDS:
            good = [q.strip() for q in get(f"{d}__good_evidence").split(" || ") if q.strip()]
            bad = [q.strip() for q in get(f"{d}__bad_evidence").split(" || ") if q.strip()]
            if good or bad:
                evidence[d] = {"good": good, "bad": bad}
        labels.append({
            "labeler": get("labeler"),
            "source_row": norm_row(r[hdr["source_row"]]),
            "ceo_id": get("ceo_id"),
            "would_recommend": get("would_recommend") or None,
            "overall_notes": get("overall_notes") or None,
            "scores": scores,
            "notes": notes,
            "evidence": evidence,
        })
    return labels


def main():
    if db.IS_PG:
        # Seeding wipes everything, including any round-2 scores already entered.
        # Never let that happen to the live DB without an explicit confirmation.
        existing = db._query_one(
            "SELECT COUNT(*) AS n FROM assignments WHERE status != 'pending'"
        ) if _assignments_table_exists() else None
        n_scored = existing["n"] if existing else 0
        print(f"Target: REMOTE Postgres ({n_scored} scored assignment(s) present).")
        answer = input("This WIPES all transcripts, prior labels, and round-2 scores "
                       "in that database. Type WIPE to continue: ").strip()
        if answer != "WIPE":
            print("Aborted — nothing changed.")
            return
    db.init_db()
    transcripts = read_transcripts()
    labels = read_labels()
    wanted_rows = sorted(transcripts, key=lambda s: int(s) if s.isdigit() else 0)
    orphan = sorted({l["source_row"] for l in labels} - set(wanted_rows))
    if orphan:
        print(f"WARNING: {len(orphan)} labeled source_rows not in transcripts file: {orphan}")

    conn = db.connect()
    try:
        ts = db.now()
        for table in ("assignments", "prior_labels", "items", "labelers"):
            conn.execute(f"DELETE FROM {table}")

        item_ids = {}
        for src in wanted_rows:
            if src not in transcripts:
                continue
            t = transcripts[src]
            conn.execute(
                db.Q("INSERT INTO items (source_row, ceo_id, transcript, created_at) VALUES (?,?,?,?)"),
                (src, t["ceo_id"], t["transcript"], ts),
            )
        for r in conn.execute("SELECT id, source_row FROM items").fetchall():
            item_ids[r["source_row"]] = r["id"]

        n_prior = 0
        for l in labels:
            iid = item_ids.get(l["source_row"])
            if not iid:
                continue
            conn.execute(
                db.Q("""INSERT INTO prior_labels
                        (item_id, labeler_name, grp, would_recommend, overall_notes,
                         scores_json, notes_json, evidence_json)
                        VALUES (?,?,?,?,?,?,?,?)"""),
                (iid, l["labeler"], group_of(l["labeler"]), l["would_recommend"],
                 l["overall_notes"], json.dumps(l["scores"]), json.dumps(l["notes"]),
                 json.dumps(l["evidence"])),
            )
            n_prior += 1

        for name, code in config.LABELERS:
            conn.execute(
                db.Q("INSERT INTO labelers (name, passcode, created_at) VALUES (?,?,?)"),
                (name, code, ts),
            )
        labeler_ids = [r["id"] for r in conn.execute("SELECT id FROM labelers").fetchall()]
        for iid in item_ids.values():
            for lid in labeler_ids:
                conn.execute(
                    db.Q("INSERT INTO assignments (item_id, labeler_id, updated_at) VALUES (?,?,?)"),
                    (iid, lid, ts),
                )
        conn.commit()
    finally:
        conn.close()

    print(f"Loaded {len(item_ids)} transcripts, {n_prior} round-1 labels, "
          f"{len(config.LABELERS)} labelers x {len(item_ids)} assignments.")
    for name, code in config.LABELERS:
        print(f"  {name}: passcode {code}  (magic link: /l/{code})")


if __name__ == "__main__":
    main()
