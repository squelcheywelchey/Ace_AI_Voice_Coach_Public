"""Generate 'Makes It a Dialogue - Judge Calls to Be Made + Takeaways.docx',
matching the format/palette of the other per-dimension takeaways docs."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x17, 0x36, 0x5D)
BLUE = RGBColor(0x36, 0x5F, 0x91)
MUTED = RGBColor(0x59, 0x59, 0x59)
INK = RGBColor(0x21, 0x21, 0x21)
HEAD_FONT, BODY_FONT = "Cambria", "Calibri"

doc = Document()

# page: US Letter + comfortable margins
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.75)
sec.left_margin = sec.right_margin = Inches(0.9)

# base body font
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = INK


def runs(par, parts, *, font=BODY_FONT, size=11, color=INK, italic=False):
    """parts: list of (text, bold) tuples or a plain string."""
    if isinstance(parts, str):
        parts = [(parts, False)]
    for text, bold in parts:
        r = par.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.italic = italic


def title(text):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(4)
    runs(p, [(text, True)], font=HEAD_FONT, size=20, color=NAVY)


def meta(parts):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    runs(p, parts, size=9, color=MUTED)


def h1(text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(4)
    runs(p, [(text, True)], font=HEAD_FONT, size=14, color=BLUE)


def body(parts, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    runs(p, parts)


def bullet(parts):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    runs(p, parts)


B = lambda t: (t, True)
T = lambda t: (t, False)

# ---------------------------------------------------------------- content
title("Makes It a Dialogue — Judge Calls to Be Made + Takeaways")
meta([B("Labels: "),
      T("SYNTHETIC seed set — 18 author-written transcripts, human-reviewed by Maya (2026-07-21); consensus in judge-suite/labels/consensus_makes_it_a_dialogue.csv. "),
      B("Judge: "), T("makes_it_a_dialogue_v1.yaml (Claude, sonnet-4-6, temperature 0).")])
meta([B("As of: "), T("2026-07-21. "), B("Status: "),
      T("v1 SHIPPED to production (9th live judge). Labels reviewed and confirmed by Maya (2026-07-21); calibration is still SYNTHETIC-ONLY and single-reviewer — independent multi-grader sign-off on the pass class is the open call.")])

h1("Current score")
body([B("Dev: "), T("100% (16/16). TPR 1.00 · TNR 1.00 · PPV 1.00 · NPV 1.00. Balanced 8 pass / 8 fail, including 5 adversarial positives and 5 adversarial negatives.")])
body([B("Test (held-out, one-shot): "), T("100% (2/2). One pass, one fail.")])
body([B("Real-corpus spot-check: "), T("12/12 real transcripts correctly FAIL (0 false-fires), including row P-755 — the rubric's named nearest-miss content follow-up.")])
body([B("Read: "), T("v1 is internally airtight on the synthetic set and holds the FAIL line on real calls. The labels are human-confirmed — Maya reviewed all 18 on 2026-07-21 — but the positives are synthetic and were confirmed by a single reviewer who co-authored the set, not the independent multi-grader consensus used for the other dimensions. So the 100% shows the judge agrees with the reviewed labels, not that it matches an independent human panel or a real coach producing dialogue (none does yet). Treat it as a robust, reviewed prompt — not yet a corpus-calibrated metric.")])

h1("Confident that")
bullet([B("No coach in the corpus meets the pass bar. "), T("The pass class defines TARGET behavior; every real session fails today. This is one face of the core “coach lectures instead of running a practice loop” failure mode.")])
bullet([B("The dimension turns on ONE distinction: "), T("a reflective question (turns attention on the participant's OWN answer/experience — “how did that feel?”, “how do you think that went?”) vs a content follow-up (probes the substance — “say more about X”, “what would you add?”). Only the former counts.")])
bullet([B("Both conditions are required and enforced: "), T("a reflective question AND engagement with the participant's specifics. Reflective questions with no specific engagement FAIL — v1 catches this (row 9107).")])
bullet([B("v1 recognizes reflective questions by category, not keyword: "), T("it passes non-canonical phrasings (“if you played that back, what would you tighten?”, “how do you feel you did today?”) and one buried in disfluent STT.")])
bullet([B("It rejects look-alikes: "), T("comprehension checks (“does that make sense?”), factual clarifiers (“was that at your last job?”), and content-add questions all FAIL — even under STT noise.")])

h1("Current rule (operative pass/fail logic — v1 as shipped)")
body([B("PASS "), T("if, anywhere in the session, the coach (AI: turns) asks at least one reflective question that turns attention back on the participant's own answer / performance / experience, AND the coach's feedback engages the specifics of what the participant actually said. One reflective question anywhere is enough; it need not recur.")])
body([B("FAIL "), T("if no reflective question appears (however rich the content engagement), if feedback is one-way templates that never engage specifics, or if the coaching is ONLY reflective questions with no substantive engagement.")])
body([B("N/A "), T("only if the call ends before any coaching feedback occurs (truncated / verification-failure calls). Essentially never on a completed session.")])

h1("Decisions made (this round)")
body([B("1. Pass bar unchanged from Rubric v2 DECIDED. "), T("“At least one reflective question anywhere + engagement with specifics” (criterion adopted from Grader A + CEO program lead). Faithful transcription; not relaxed.")])
body([B("2. Synthetic seeding is legitimate here. "), T("With zero real positives, the pass class must be authored to exist at all — same posture as pii. Maya reviewed all 18 labels (2026-07-21) and confirmed them.")])
body([B("3. Reflective vs content follow-up is the operative boundary. "), T("“What would you tighten?” (self-evaluative) PASSES (row 9007); “what else could you add?” (content generation) FAILS (row 9108). Comprehension checks and factual clarifiers FAIL.")])
body([B("4. Shipped v1 to production despite synthetic-only calibration, "), T("with the pii-style caveat: FAIL-dominant on live calls; any PASS is a review flag, not ground truth (migration 005, commit 2c80e91).")])

h1("Decisions needed (team)")
body([B("A. Independent sign-off on the pass class. "), T("Maya has reviewed and confirmed all 18 labels (2026-07-21), so they are human-confirmed — but by a single reviewer who co-authored the set, on synthetic transcripts. The other dimensions were calibrated to independent multi-grader consensus; to meet that bar, Grader B / Grader C should independently adjudicate the 9 positives (especially the tricky ones) before this judge's PASS is trusted as a metric. Until then, PASS is confirmed-but-not-independently-calibrated. Decision: ____")])
body([B("B. Is “one reflective question anywhere” the right bar? "), T("Row 9006 passes on a single reflective question amid an otherwise templated coach. Is that genuinely dialogue, or should the bar require more (e.g., reflection that actually feeds the next turn)? Decision: ____")])
body([B("C. The tighten-vs-add boundary (9007 PASS ↔ 9108 FAIL). "), T("Is “what would you tighten?” really reflective while “what would you add?” is not? If the team rejects the distinction, both rows move together. Decision: ____")])
body([B("D. First-live-PASS protocol. "), T("Since a live PASS is a real event (coach hit target behavior OR judge false-fire), agree to pull and human-review the first N live PASS transcripts — the real-world validation the synthetic set cannot provide. Decision: ____")])

h1("Key lessons")
bullet([B("A dimension with no real positives can still ship — but only as instrumented discovery. "), T("Synthetic seeding lets the judge exist and be pressure-tested; it cannot substitute for human agreement.")])
bullet([B("Acing a synthetic set proves prompt soundness, not calibration. "), T("16/16 says the definition is internally consistent and robust to adversarial cases; it says nothing about whether humans would draw the line where Claude did.")])
bullet([B("Matched contrastive pairs are the highest-value test data. "), T("9006↔9107 and 9007↔9108 differ only on the exact variable the dimension turns on; a judge that splits them correctly has learned the boundary, not the surface.")])
bullet([B("The real signal is on real calls. "), T("Holding the FAIL line on 12 real transcripts (0 false-fires) — including the rubric's named nearest miss — is stronger evidence than any synthetic score, and marks where the first genuine PASS should be scrutinized.")])
bullet([B("This is the “no practice loop” failure mode wearing a judge. "), T("A live PASS rate above ~0 is itself a coach-improvement signal, independent of per-call calibration.")])

h1("Run history")
body([B("v1 (makes_it_a_dialogue_v1.yaml) "), T("— faithful transcription of Rubric v2 DECIDED; bright line built on the reflective-vs-content-followup distinction. Dev 100% (16/16), TPR 1.00 / TNR 1.00. Test 100% (2/2). Real-corpus spot-check 12/12 FAIL, 0 false-fires. No iteration needed — v1 held across easy and adversarial cases in both directions.")])
body([B("Production "), T("— shipped 2026-07-21 as the 9th live judge (ceo-eval-server, migration 005, commit 2c80e91). Synthetic-only caveat carried in the migration comment.")])
body([B("Audit trail: "), T("prompt (quotes no real transcript text — no leakage), per-run metrics (judge-suite/iteration_log/run_log.csv), sha-pinned snapshot (makes_it_a_dialogue_v1_668e72a0.yaml), the 18-row set + Maya's label-review page (build_dialogue_review.py), and the iteration log (judge-suite/iteration_log/makes_it_a_dialogue.md).")], after=2)

out = "docs/judges/Makes It a Dialogue - Judge Calls to Be Made + Takeaways.docx"
doc.save(out)
print("wrote", out)
