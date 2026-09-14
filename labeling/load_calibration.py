"""Load a calibration round from an assignments spreadsheet.

The spreadsheet has one row per (transcript, grader) assignment with columns:
    row_idx | ceo_id | excerpt | grader_name

  * row_idx     - stable transcript number shown to graders ("Transcript N")
  * ceo_id      - participant id (kept for admin/export; hidden from graders)
  * excerpt     - the transcript text the grader scores
  * grader_name - the annotator the row is assigned to

(Older headers CEO ID / Annotator / Transcript / Index are accepted as aliases.)

Loading (re)builds the round: an account per unique grader, one transcript per
unique ceo_id, and the assignments connecting them. A transcript listed under
two graders is imported once and assigned to both (double-rating). Loading is
additive and idempotent — existing rows are reused, not duplicated.

Used two ways:
  * CLI:  python load_calibration.py <file.xlsx> [--fresh]
  * Web:  the admin "Load calibration round" upload (calls load_assignments).
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

import config
import db

# header -> accepted aliases (lowercased)
COL_ALIASES = {
    "ceo": ["ceo_id", "ceo id"],
    "grader": ["grader_name", "annotator", "grader"],
    "excerpt": ["excerpt", "transcript"],
    "row_idx": ["row_idx", "index", "idx", "#"],
}


def load_assignments(ws, source_name: str) -> dict:
    """Create fellows/transcripts/assignments from a worksheet.

    Returns counts: {transcripts, fellows, assignments, new_fellows:[(name,code)]}.
    Raises ValueError if the expected columns aren't present.
    """
    header = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]

    def find_col(key):
        for alias in COL_ALIASES[key]:
            if alias in header:
                return header.index(alias) + 1
        return None

    ceo_c, grader_c, exc_c = find_col("ceo"), find_col("grader"), find_col("excerpt")
    idx_c = find_col("row_idx")  # optional
    if not (ceo_c and grader_c and exc_c):
        raise ValueError(
            f"Expected columns row_idx / ceo_id / excerpt / grader_name; got {header}"
        )

    rows = []
    for r in range(2, ws.max_row + 1):
        ceo = ws.cell(r, ceo_c).value
        annot = ws.cell(r, grader_c).value
        txt = ws.cell(r, exc_c).value
        if ceo is None or not annot or not txt:
            continue
        index = ws.cell(r, idx_c).value if idx_c else None
        rows.append((str(ceo).strip(), str(annot).strip(), str(txt), index))

    db.ensure_admin(config.ADMIN_PASSCODE)

    # graders not already in the DB get a fresh passcode (names are unique here)
    existing = {l["name"] for l in db.all_labelers()}
    new_fellows, seen = [], set()
    for _, annot, _, _ in rows:
        if annot not in existing and annot not in seen:
            seen.add(annot)
            new_fellows.append((annot, db.gen_passcode(annot)))

    # one transcript per unique CEO ID; its number is the uploaded row_idx
    # (fallback to a 1..N counter), used as the item's source_row.
    item_specs, ceo_to_src, counter = [], {}, 0
    for ceo in dict.fromkeys(c for c, _, _, _ in rows):
        counter += 1
        rec = next(x for x in rows if x[0] == ceo)
        try:
            number = int(rec[3])
        except (TypeError, ValueError):
            number = counter
        ceo_to_src[ceo] = number
        item_specs.append((number, ceo, rec[2]))

    assignment_pairs = [(ceo_to_src[ceo], annot) for ceo, annot, _, _ in rows]

    # one connection for the whole load (fast over remote Postgres)
    db.import_round(source_name, new_fellows, item_specs, assignment_pairs)

    return {
        "transcripts": len(item_specs),
        "fellows": len({a for _, a, _, _ in rows}),
        "assignments": len(set(assignment_pairs)),
        "new_fellows": new_fellows,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file", help="assignments .xlsx (CEO ID | Annotator | Transcript)")
    ap.add_argument("--fresh", action="store_true",
                    help="back up and wipe the existing DB before loading")
    args = ap.parse_args()

    path = Path(args.file)
    if not path.exists():
        sys.exit(f"No such file: {path}")

    if args.fresh and config.DB_PATH.exists():
        backup = config.DB_PATH.with_suffix(".db.bak")
        config.DB_PATH.replace(backup)
        print(f"Backed up existing DB -> {backup.name}")

    db.init_db()
    ws = load_workbook(path, data_only=True).worksheets[0]
    counts = load_assignments(ws, source_name=path.name)

    for name, code in counts["new_fellows"]:
        print(f"  fellow: {name:10s} passcode={code}  link=/l/{code}")
    print(f"\nLoaded {counts['transcripts']} transcript(s), {counts['fellows']} "
          f"fellow(s), {counts['assignments']} assignment(s).")
    print(f"Admin passcode: {config.ADMIN_PASSCODE}")


if __name__ == "__main__":
    main()
