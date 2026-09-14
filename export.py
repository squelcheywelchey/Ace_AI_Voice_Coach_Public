"""Export coach_eval.jsonl + transcripts into a filterable spreadsheet.

Produces coach_eval.xlsx — one row per call, with the quality label, every
failure-mode flag, every rubric score, the evidence quotes, the call summary,
and the full transcript. Open it in Excel or Google Sheets and use the filter
arrows in the header row (Data > Filter) to slice by any column.

Usage:
    python export.py
    python export.py --scores coach_eval.jsonl --tsv pcpt_mock_interview_v2.tsv --out coach_eval.xlsx
"""

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analyze import RUBRIC, FLAGS, SECTIONS, GROWTH_DIMS, BASELINE_DIMS, _mean_scores, value_add

csv.field_size_limit(10**7)
CRITERIA = [(sec, cid, name) for sec, cid, name, *_ in RUBRIC]
FLAG_IDS = [fid for fid, _ in FLAGS]


def load_scores(path: Path) -> dict:
    by_row = {}
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                r = json.loads(line)
                by_row[r["row"]] = r
    return by_row


def load_transcripts(path: Path) -> dict:
    by_row = {}
    with open(path, newline="", encoding="utf-8") as fh:
        for i, row in enumerate(csv.DictReader(fh, delimiter="\t")):
            by_row[i] = row
    return by_row


def avg_score(r: dict):
    vals = [int(r[cid]) for _s, cid, _n in CRITERIA if r.get(cid) not in (None, "N/A")]
    return round(sum(vals) / len(vals), 2) if vals else ""


def main() -> None:
    p = argparse.ArgumentParser(description="Export scores + transcripts to a filterable .xlsx")
    p.add_argument("--scores", default="coach_eval.jsonl")
    p.add_argument("--tsv", default="pcpt_mock_interview_v2.tsv")
    p.add_argument("--out", default="coach_eval.xlsx")
    args = p.parse_args()

    scores = load_scores(Path(args.scores))
    tsv = load_transcripts(Path(args.tsv))

    # Column layout
    key_cols = ["row", "ceo_id", "call_minutes", "overall_quality", "avg_score",
                "coach_growth", "participant_baseline", "value_add"]
    flag_cols = list(FLAG_IDS)
    score_cols = [cid for _s, cid, _n in CRITERIA]
    text_cols = ["summary", "flag_evidence", "rubric_red_flags", "transcript"]
    headers = key_cols + flag_cols + score_cols + text_cols

    wb = Workbook()
    ws = wb.active
    ws.title = "coach scores"

    # Header row
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="DDDDDD")
    for c, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=False)

    for row_idx in sorted(scores):
        r = scores[row_idx]
        src = tsv.get(row_idx, {})
        fe = "; ".join(f"{x['flag']}: {x['quote']}" for x in r.get("flag_evidence", []))
        rf = "; ".join(f"{x['criterion']}: {x['quote']}" for x in r.get("red_flags", []))
        line = {
            "row": row_idx,
            "ceo_id": r.get("ceo_id"),
            "call_minutes": src.get("Call Duration (minutes with decimal precision)", ""),
            "overall_quality": r.get("overall_quality"),
            "avg_score": avg_score(r),
            "coach_growth": (round(_mean_scores(r, GROWTH_DIMS), 2) if _mean_scores(r, GROWTH_DIMS) is not None else ""),
            "participant_baseline": (round(_mean_scores(r, BASELINE_DIMS), 2) if _mean_scores(r, BASELINE_DIMS) is not None else ""),
            "value_add": (value_add(r) if value_add(r) is not None else ""),
            **{fid: bool(r.get(fid)) for fid in flag_cols},
            **{cid: (int(r[cid]) if r.get(cid) not in (None, "N/A") else "N/A") for cid in score_cols},
            "summary": r.get("summary", ""),
            "flag_evidence": fe,
            "rubric_red_flags": rf,
            "transcript": src.get("Transcript", ""),
        }
        ws.append([line[h] for h in headers])

    # Filter + freeze the key columns and header row
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = f"{get_column_letter(len(key_cols) + 1)}2"  # freeze key cols + header

    # Column widths
    widths = {"row": 6, "ceo_id": 10, "call_minutes": 8, "overall_quality": 13, "avg_score": 9,
              "coach_growth": 12, "participant_baseline": 18, "value_add": 10,
              "summary": 60, "flag_evidence": 70, "rubric_red_flags": 70, "transcript": 120}
    for c, name in enumerate(headers, 1):
        letter = get_column_letter(c)
        if name in widths:
            ws.column_dimensions[letter].width = widths[name]
        elif name in flag_cols:
            ws.column_dimensions[letter].width = max(12, len(name) // 2)
        else:  # score columns — keep narrow, name is long so rely on filter dropdown
            ws.column_dimensions[letter].width = 7

    wb.save(args.out)
    print(f"Wrote {args.out} — {len(scores)} calls, {len(headers)} columns.")
    print("Open it and click the filter arrows in row 1 (Data > Filter if not already on).")
    print("Score columns use the short ids; see findings.md / summary_report.txt for full names.")


if __name__ == "__main__":
    main()
